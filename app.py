from pathlib import Path
import io, re, shutil, tempfile, zipfile
import streamlit as st
from docx import Document

BASE=Path(__file__).parent
T=BASE/'templates'

CASI_TD=[
'1 - Particolarità tecnologiche/infrastrutturali / continuità',
'2 - Emergenza adeguatamente motivata','3 - Consulenze/prestazioni professionali specialistiche o complementari',
'4 - Acquisto di opportunità','5 - Gara andata deserta','6 - Prezzi imposti / cessione azienda / impossibilità oggettiva di competizione',
'7 - Estensione condizioni/prezzi di competizione','8 - Progetto business/tecnologico identificato con fabbisogni e fornitori puntuali',
'9 - Commercializzazione prodotti/soluzioni non standard a clienti privati','10 - Opportunità business per commercializzazione verso PA']

SPECIALI=[
'1 - Scelte tecnologiche/business di Gruppo da unico fornitore','2 - Diritti/sponsorizzazioni/collaboratori artistici unici',
'3 - Manutenzione applicativi proprietari / rete non TPM','4 - Partnership commerciali','5 - Canali del Costruttore certificati / prezzi Procurement',
'6 - Vendor tecnologico/distributore per gare PA con tecnologia esplicita ed esclusività','7 - Nuove tecnologie / primo impiego / sperimentazioni',
'8 - Postazioni alta frequenza','9 - Prodotti di rete per ampliamento reti esistenti','10 - Adesione ad Accordi Quadro di Gruppo',
'11 - Networking ICT con accordi/listini PR >=80% commessa','12 - Leasing individuato da Enterprise/CF.F']

def euro(s):
    if isinstance(s,(int,float)): return float(s)
    s=str(s).strip().replace('€','').replace(' ','')
    if ',' in s: s=s.replace('.','').replace(',','.')
    return float(s or 0)

def fmt_eur(x):
    return f"€ {x:,.2f}".replace(',', 'X').replace('.', ',').replace('X','.')

def decide(d):
    docs=[]; reasons=[]; warnings=[]
    # Nuova procedura: il motore evita automatismi non supportati e chiede le condizioni rilevanti.
    docs += ['4','6']
    reasons += ['Allegato 4: requisiti sicurezza/compliance','Allegato 6: valutazione trattamento dati e firma del modulo']
    if not d['nuova_tranche_senza_variazione'] and not d['precontrattuale']:
        docs.append('7'); reasons.append('Allegato 7: razionali fabbisogno/prezzo')
    if d['deroga'] or d['side_letter']:
        docs.append('OFFERTA'); reasons.append('Offerta: acquisto in deroga o gruppo merce Side Letter')
    if d['trattativa_diretta'] and d['totale']>20000 and not d['acquisto_speciale'] and not d['infragruppo'] and not d['accordo_esistente']:
        docs.append('3A'); reasons.append('Allegato 3A: TD >20k e nessuna esclusione dichiarata')
    if d['gara_prest_prof']:
        docs.append('13'); reasons.append('Allegato 13: gara per prestazioni professionali')
    if d['saas']:
        docs.append('14'); reasons.append('Allegato 14: checklist sicurezza SaaS')
    if d['cliente_tipo']=='Privato' and d['trattamento_cliente']:
        docs.append('11'); reasons.append('Allegato 11: ATCS cliente privato')
    if d['cliente_tipo']=='PA' and d['trattamento_cliente']:
        docs.append('12'); reasons.append('Allegato 12: ATCS cliente PA')
    if d['amministratore_sistema']:
        docs.append('15'); reasons.append('Allegato 15: attività di Amministratore di Sistema')
    if d['trattamento_dati']:
        docs.append('10'); reasons.append('Allegato 10: istruzioni data breach per responsabili terzi')
    if d['attestazione_conformita']:
        docs.append('16'); reasons.append('Allegato 16: attestazione conformità privacy')
    if d['acquisto_speciale']:
        warnings.append('Acquisto Speciale selezionato: non genera 3A automaticamente; si applica per il resto il processo standard.')
    if d['trattativa_diretta'] and not d['casistica_td'] and not d['acquisto_speciale']:
        warnings.append('Selezionare la casistica di Trattativa Diretta (Allegato 1).')
    return list(dict.fromkeys(docs)), reasons, warnings

def xml_replace(src,dst,repls,occurrence_repls=None):
    """Sostituzione sicura nei nodi di testo OOXML: non corrompe il DOCX con &, <, >, accenti."""
    from lxml import etree
    occurrence_repls=occurrence_repls or []
    counts={}
    with zipfile.ZipFile(src,'r') as zin, zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data=zin.read(item.filename)
            if item.filename.endswith('.xml'):
                try:
                    root=etree.fromstring(data)
                    for node in root.xpath('//*[local-name()="t"]'):
                        if node.text is None: continue
                        text=node.text
                        for a,b in repls.items():
                            if a in text:
                                text=text.replace(a,str(b))
                        for old,new,n in occurrence_repls:
                            if old in text:
                                counts[old]=counts.get(old,0)+1
                                if counts[old]==n:
                                    text=text.replace(old,str(new),1)
                        node.text=text
                    data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
                except Exception:
                    pass
            zout.writestr(item,data)

def make7(d,out):
    repl={
      '21/09/2026':d['data'], 'OLIVETTI - P&amp;DI CPP':d['funzione'], 'OLIVETTI - P&DI CPP':d['funzione'],
      'ANTONIO SCUCCIMARRA':d['fornitore'],
      'Nota Spese Giugno, CB Centro Sud Puglia - Acque del Sud':d['oggetto'],
      'supporto per attività commerciale propedeutica alla riuscita del Progetto Water Management System gestito da Portfolio IoT':d['motivazione'],
      '216,40':fmt_eur(d['totale']).replace('€ ','')
    }
    xml_replace(T/'allegato7.docx',out,repl)

def make6(d,out):
    repl={'2101416996':d['rda']}
    # Il modello caricato contiene due placeholder consecutivi nella riga Riferimenti/Fornitore.
    occ=[('Scrivere qui',d['riferimenti'],1),('Scrivere qui',d['fornitore'],2)]
    xml_replace(T/'allegato6.docx',out,repl,occ)

def make3a(d,out):
    # Lavora sul DOCX originale: niente sostituzioni globali di numeri/placeholder.
    shutil.copy2(T/'allegato3.docx',out)
    doc=Document(out)
    # Data del solo 3A (non altera i placeholder del 3B).
    for p in doc.paragraphs:
        if p.text.strip()=='Luogo, gg/mm/aaaa':
            p.text='Luogo Roma, '+d.get('data','')
            break
    t=doc.tables[1]
    t.cell(2,0).text='Fornitore proposto:\n'+d.get('fornitore','')
    t.cell(2,1).text='Cod. Sap: '+d.get('sap','')
    t.cell(2,3).text='P.IVA: '+d.get('piva','')
    t.cell(3,0).text='Importo TD: '+fmt_eur(d.get('totale',0))
    if d.get('importo_cumulato'):
        t.cell(3,1).text='Importo cumulato dei BO/contratti comprensivo della TD in oggetto:\n'+fmt_eur(d['importo_cumulato'])
    t.cell(3,3).text='N° RDA (se emessa): '+d.get('rda','')
    if d.get('tipo_td'): t.cell(4,0).text='La richiesta di Trattativa Diretta è per: '+d['tipo_td']
    if d.get('ultimo_contratto'): t.cell(4,1).text='Numero ultimo contratto/BO di riferimento:\nN.°: '+d['ultimo_contratto']
    if d.get('data_inizio') or d.get('data_fine'): t.cell(4,3).text=f"Periodo fornitura:\ninizio: {d.get('data_inizio','')}    fine: {d.get('data_fine','')}"
    if d.get('pluriennale'): t.cell(5,0).text='Attività pluriennale (progetto di durata superiore ad 1 anno): '+d['pluriennale']
    if d.get('casistica_td'): t.cell(6,0).text="Valorizzare il campo con una delle casistiche dell’allegato 1: "+d['casistica_td']
    if d.get('descrizione_3a'): t.cell(7,0).text='Descrizione chiara, dettagliata e circostanziata dell’oggetto della richiesta, con indicazione dei tempi di realizzazione:\n'+d['descrizione_3a']
    if d.get('acquisto_tecnologia'): t.cell(8,0).text='Acquisto di tecnologia: '+d['acquisto_tecnologia']
    if d.get('tecnologia'): t.cell(8,2).text='Indicare la tecnologia:\n'+d['tecnologia']
    if d.get('vincolo_tecnologico'): t.cell(9,0).text='Vincolo tecnologico / legame fornitore-cliente-vendor:\n'+d['vincolo_tecnologico']
    if d.get('rischi_altro_fornitore'): t.cell(10,0).text='Rischi derivanti dall’affidamento a fornitore diverso da quello proposto:\n'+d['rischi_altro_fornitore']
    if d.get('legacy'): t.cell(11,0).text='Continuità tecnologica / grado di legacy e tempi per avvio processo competitivo:\n'+d['legacy']
    if d.get('continuita'): t.cell(12,0).text='La richiesta è relativa ad attività in continuità con lo stesso fornitore? '+d['continuita']
    if d.get('fornitore_impegnato'): t.cell(13,0).text='Fornitore è già impegnato? '+d['fornitore_impegnato']
    if d.get('stato_fornitura'): t.cell(14,0).text='Stato della fornitura: '+d['stato_fornitura']
    if d.get('motivazione_sanatoria'): t.cell(15,0).text='Motivazione esclusione Atto a Sanatoria:\n'+d['motivazione_sanatoria']
    if d.get('documenti_allegati'): t.cell(16,0).text='Elenco documenti allegati:\n'+d['documenti_allegati']
    if d.get('cliente'): t.cell(17,0).text='Cliente destinatario della fornitura: '+d['cliente']
    doc.save(out)

def package(d,docs,offer_bytes=None,offer_name=None):
    td=Path(tempfile.mkdtemp())
    produced=[]
    if '4' in docs: shutil.copy2(T/'allegato4.docx',td/f"RDA {d['rda']} - Allegato 4.docx"); produced.append(td/f"RDA {d['rda']} - Allegato 4.docx")
    if '6' in docs: make6(d,td/f"RDA {d['rda']} - Allegato 6.docx"); produced.append(td/f"RDA {d['rda']} - Allegato 6.docx")
    if '7' in docs: make7(d,td/f"RDA {d['rda']} - Allegato 7.docx"); produced.append(td/f"RDA {d['rda']} - Allegato 7.docx")
    if '3A' in docs: make3a(d,td/f"RDA {d['rda']} - Allegato 3A - DA VERIFICARE.docx"); produced.append(td/f"RDA {d['rda']} - Allegato 3A - DA VERIFICARE.docx")
    for n in ['10','11','12','13','15','16']:
        if n in docs:
            shutil.copy2(T/f'allegato{n}.docx',td/f"RDA {d['rda']} - Allegato {n}.docx"); produced.append(td/f"RDA {d['rda']} - Allegato {n}.docx")
    if '14' in docs:
        shutil.copy2(T/'allegato14.xlsx',td/f"RDA {d['rda']} - Allegato 14.xlsx"); produced.append(td/f"RDA {d['rda']} - Allegato 14.xlsx")
    if 'OFFERTA' in docs:
        if offer_bytes:
            p=td/(offer_name or 'Offerta.pdf'); p.write_bytes(offer_bytes); produced.append(p)
        else:
            p=td/'OFFERTA_MANCANTE.txt'; p.write_text('Caricare l’offerta del fornitore prima di chiudere la RDA.',encoding='utf-8'); produced.append(p)
    z=io.BytesIO()
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as zz:
        for p in produced: zz.write(p,p.name)
    return z.getvalue()

st.set_page_config(page_title='Generatore RDA Olivetti v0.4.2',layout='wide')
st.title('Generatore RDA Olivetti — v0.4.2')
st.caption('Nuova procedura: motore decisionale + compilazione sui DOCX originali. I campi non supportati non vengono inventati.')

c1,c2,c3=st.columns(3)
with c1:
    rda=st.text_input('N. RDA')
    fornitore=st.text_input('Fornitore')
    sap=st.text_input('Codice SAP fornitore')
    piva=st.text_input('P. IVA fornitore')
    totale=euro(st.text_input('Totale RDA','0,00'))
with c2:
    data=st.text_input('Data','21/09/2026')
    funzione=st.text_input('Funzione richiedente','OLIVETTI - P&DI CPP')
    cliente=st.text_input('Cliente / Gara')
    cliente_tipo=st.selectbox('Tipo cliente',['n.a.','Privato','PA'])
    riferimenti=st.text_input('Riferimenti (SDW/SH/PCD, se applicabile)')
with c3:
    oggetto=st.text_area('Oggetto')
    motivazione=st.text_area('Motivazione / descrizione esigenza')

st.subheader('Domande decisionali')
a,b,c=st.columns(3)
with a:
    pre_sales=st.checkbox('RDA Pre Sales')
    trattativa_diretta=st.checkbox('Trattativa Diretta')
    acquisto_speciale=st.checkbox('Acquisto Speciale (Allegato 2)')
    infragruppo=st.checkbox('Acquisto infragruppo')
with b:
    accordo_esistente=st.checkbox('Contratto/AQ/listino già in essere')
    nuova_tranche_senza_variazione=st.checkbox('Nuova tranche senza variazione valore complessivo')
    precontrattuale=st.checkbox('Ingaggio fornitore in fase precontrattuale')
    deroga=st.checkbox('Acquisto in deroga')
    side_letter=st.checkbox('Gruppo merce soggetto a Side Letter')
with c:
    trattamento_dati=st.checkbox('La fornitura tratta dati personali')
    trattamento_cliente=st.checkbox('Trattamento dati di cui il cliente è Titolare/Responsabile')
    saas=st.checkbox('Servizio SaaS')
    amministratore_sistema=st.checkbox('Attività di Amministratore di Sistema')
    gara_prest_prof=st.checkbox('Gara per prestazioni professionali')
    attestazione_conformita=st.checkbox('Richiesta attestazione conformità privacy')

casistica_td=''
if trattativa_diretta and not acquisto_speciale:
    casistica_td=st.selectbox('Casistica Trattativa Diretta — Allegato 1',['']+CASI_TD)
if acquisto_speciale:
    st.selectbox('Casistica Acquisto Speciale — Allegato 2',['']+SPECIALI)

# Sezione 3A sempre visibile nel flusso quando la pratica è una TD.
# Il motore mostra anche il motivo per cui il documento sarà o non sarà generato.
st.subheader('Allegato 3A — verifica e compilazione')
three_a_eligible = trattativa_diretta and totale>20000 and not acquisto_speciale and not infragruppo and not accordo_esistente
if not trattativa_diretta:
    st.info('Allegato 3A: NO — la pratica non è stata indicata come Trattativa Diretta.')
elif totale<=20000:
    st.info('Allegato 3A: NO — importo non superiore a € 20.000.')
elif acquisto_speciale:
    st.info('Allegato 3A: NO — è stato selezionato Acquisto Speciale.')
elif infragruppo:
    st.info('Allegato 3A: NO — è stato selezionato Acquisto infragruppo.')
elif accordo_esistente:
    st.info('Allegato 3A: NO — è stato indicato un contratto/AQ/listino già in essere.')
else:
    st.success('Allegato 3A: SÌ — verrà inserito nel pacchetto e precompilato con i dati sottostanti.')

importo_cumulato=0.0; tipo_td=''; ultimo_contratto=''; data_inizio=''; data_fine=''; pluriennale=''
descrizione_3a=''; acquisto_tecnologia=''; tecnologia=''; vincolo_tecnologico=''; rischi_altro_fornitore=''
legacy=''; continuita=''; fornitore_impegnato=''; stato_fornitura=''; motivazione_sanatoria=''; documenti_allegati=''
if three_a_eligible:
    x1,x2=st.columns(2)
    with x1:
        importo_cumulato=euro(st.text_input('3A — Importo cumulato BO/contratti comprensivo TD','0,00'))
        tipo_td=st.selectbox('3A — La richiesta di TD è per',['','Nuova fornitura','Rinnovo','Estensione/variante','Altro'])
        ultimo_contratto=st.text_input('3A — Numero ultimo contratto/BO di riferimento')
        data_inizio=st.text_input('3A — Inizio fornitura (gg/mm/aaaa)')
        data_fine=st.text_input('3A — Fine fornitura (gg/mm/aaaa)')
        pluriennale=st.selectbox('3A — Attività pluriennale',['','No','Sì'])
        acquisto_tecnologia=st.selectbox('3A — Acquisto di tecnologia',['','No','Sì'])
        tecnologia=st.text_input('3A — Tecnologia, se applicabile')
    with x2:
        descrizione_3a=st.text_area('3A — Descrizione dettagliata',value=motivazione)
        vincolo_tecnologico=st.text_area('3A — Vincolo tecnologico / vendor, se applicabile')
        rischi_altro_fornitore=st.text_area('3A — Rischi con un fornitore diverso')
        legacy=st.text_area('3A — Legacy / continuità tecnologica, se applicabile')
        continuita=st.selectbox('3A — Continuità con lo stesso fornitore',['','No','Sì'])
        fornitore_impegnato=st.selectbox('3A — Fornitore già impegnato',['','No','Sì'])
        stato_fornitura=st.selectbox('3A — Stato fornitura',['','Non avviata','In corso','Completata'])
        motivazione_sanatoria=st.text_area('3A — Motivazione esclusione Atto a Sanatoria, se applicabile')
        documenti_allegati=st.text_area('3A — Elenco documenti allegati')

d=locals().copy()
docs,reasons,warnings=decide(d)
st.subheader('Anteprima allegati')
st.write('**Da predisporre:** '+(', '.join(docs) if docs else 'nessuno'))
for x in reasons: st.write('• '+x)
for w in warnings: st.warning(w)

offer=st.file_uploader('Offerta del fornitore (se richiesta)',type=['pdf','docx','xlsx'])
if st.button('GENERA PACCHETTO RDA',type='primary',disabled=not bool(rda and fornitore and oggetto)):
    z=package(d,docs,offer.getvalue() if offer else None,offer.name if offer else None)
    st.download_button('Scarica ZIP RDA',z,file_name=f'RDA_{rda}_allegati.zip',mime='application/zip')
    if '3A' in docs: st.info('Il 3A viene precompilato nei campi sicuri e marcato DA VERIFICARE: contiene campi autorizzativi/contrattuali che non vanno inventati.')
